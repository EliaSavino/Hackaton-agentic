from __future__ import annotations

from pathlib import Path
from typing import Any

from pydantic import BaseModel, Field

from hackathon_agents.tools.base import error_result, ok_result


class ReportInput(BaseModel):
    output_path: str
    title: str = "Scientific Discovery Agent Report"
    user_request: str
    plan: dict[str, Any] | None = None
    candidates: list[dict[str, Any]] = Field(default_factory=list)
    tool_results: list[dict[str, Any]] = Field(default_factory=list)
    critic_notes: list[str] = Field(default_factory=list)
    errors: list[str] = Field(default_factory=list)


def write_scientific_report(input_data: ReportInput | dict[str, Any]):
    parsed = input_data if isinstance(input_data, ReportInput) else ReportInput.model_validate(input_data)
    try:
        from docx import Document
    except Exception:
        return error_result("python-docx is not installed or could not be imported.", {"output_path": parsed.output_path})

    try:
        output = Path(parsed.output_path)
        output.parent.mkdir(parents=True, exist_ok=True)

        document = Document()
        document.add_heading(parsed.title, level=0)
        document.add_heading("Request", level=1)
        document.add_paragraph(parsed.user_request)

        if parsed.plan:
            document.add_heading("Plan", level=1)
            for step in parsed.plan.get("steps", []):
                name = step.get("name", "step")
                description = step.get("description", "")
                document.add_paragraph(f"{name}: {description}", style="List Bullet")

        document.add_heading("Candidate Molecules", level=1)
        if parsed.candidates:
            table = document.add_table(rows=1, cols=5)
            header = table.rows[0].cells
            header[0].text = "Name"
            header[1].text = "SMILES"
            header[2].text = "Mol Wt"
            header[3].text = "LogP"
            header[4].text = "QED"
            for candidate in parsed.candidates:
                descriptors = candidate.get("descriptors", {})
                row = table.add_row().cells
                row[0].text = str(candidate.get("name") or "")
                row[1].text = str(candidate.get("smiles") or "")
                row[2].text = _format_value(descriptors.get("mol_wt"))
                row[3].text = _format_value(descriptors.get("logp"))
                row[4].text = _format_value(descriptors.get("qed"))
        else:
            document.add_paragraph("No candidate molecules were produced.")

        document.add_heading("Tool Results", level=1)
        for result in parsed.tool_results:
            tool_name = result.get("tool_name", "tool")
            ok = result.get("result", {}).get("ok")
            document.add_paragraph(f"{tool_name}: {'ok' if ok else 'failed'}", style="List Bullet")

        document.add_heading("Critic Notes", level=1)
        if parsed.critic_notes:
            for note in parsed.critic_notes:
                document.add_paragraph(note, style="List Bullet")
        else:
            document.add_paragraph("No critic notes were recorded.")

        if parsed.errors:
            document.add_heading("Errors and Fallbacks", level=1)
            for error in parsed.errors:
                document.add_paragraph(error, style="List Bullet")

        document.save(output)
        return ok_result({"path": str(output)}, [str(output)])
    except Exception as exc:
        return error_result(str(exc), {"output_path": parsed.output_path})


def _format_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        return f"{value:.3f}"
    return str(value)
