from __future__ import annotations

import json
import zipfile
from pathlib import Path
from typing import TYPE_CHECKING
from xml.sax.saxutils import escape

if TYPE_CHECKING:
    from hackathon_agents.mechanism.state import MechanismDiscoveryState


def write_mechanism_report(state: "MechanismDiscoveryState") -> tuple[Path, Path]:
    """Write JSON and DOCX reports for the current mechanism state."""

    if state.run_path is None:
        raise ValueError("state.run_dir is required before reporting")
    run_dir = state.run_path
    run_dir.mkdir(parents=True, exist_ok=True)
    json_path = run_dir / "report.json"
    docx_path = run_dir / "report.docx"
    payload = _report_payload(state)
    json_path.write_text(json.dumps(payload, indent=2), encoding="utf-8")
    _write_docx(payload, docx_path)
    state.report_json_path = str(json_path)
    state.report_docx_path = str(docx_path)
    return json_path, docx_path


def _report_payload(state: "MechanismDiscoveryState") -> dict:
    top_rankings = [ranking.model_dump(mode="json") for ranking in state.rankings[:5]]
    evidence_table = []
    for ranking in state.rankings[:5]:
        evidence_table.append(
            {
                "hypothesis_id": ranking.hypothesis_id,
                "mechanism_class": ranking.mechanism_class,
                "score": ranking.score,
                "uncertainty": ranking.uncertainty,
                "evidence": ranking.evidence,
                "caveats": ranking.caveats,
            }
        )
    return {
        "title": "Mechanism Discovery Agent Workbench",
        "objective": state.objective,
        "mode": state.mode,
        "round_index": state.round_index,
        "best_mechanism_ranking": top_rankings,
        "evidence_table": evidence_table,
        "next_recommended_experiment": state.next_experiment.model_dump(mode="json") if state.next_experiment else None,
        "uncertainty_analysis": [ranking.model_dump(mode="json") for ranking in state.rankings[:5]],
        "dft_jobs_submitted": [job.model_dump(mode="json") for job in state.dft_jobs],
        "robot_experiments_submitted": [experiment.model_dump(mode="json") for experiment in state.experiments],
        "critic_notes": state.critic_notes,
        "errors": state.errors,
        "validation_errors": state.validation_errors,
    }


def _write_docx(payload: dict, output_path: Path) -> None:
    try:
        from docx import Document
    except Exception as exc:
        _write_minimal_docx(payload, output_path, reason=str(exc))
        return

    document = Document()
    document.add_heading(payload["title"], level=0)
    document.add_heading("Objective", level=1)
    document.add_paragraph(payload["objective"])
    document.add_paragraph(f"Mode: {payload['mode']}")
    document.add_paragraph(f"Round: {payload['round_index']}")

    document.add_heading("Current Best Mechanism Ranking", level=1)
    rankings = payload["best_mechanism_ranking"]
    if rankings:
        table = document.add_table(rows=1, cols=5)
        header = table.rows[0].cells
        header[0].text = "Hypothesis"
        header[1].text = "Class"
        header[2].text = "Score"
        header[3].text = "Uncertainty"
        header[4].text = "Evidence"
        for ranking in rankings:
            row = table.add_row().cells
            row[0].text = ranking["hypothesis_id"]
            row[1].text = str(ranking["mechanism_class"])
            row[2].text = f"{ranking['score']:.3f}"
            row[3].text = f"{ranking['uncertainty']:.3f}"
            row[4].text = "; ".join(ranking.get("evidence", [])[:2])
    else:
        document.add_paragraph("No mechanism ranking is available yet.")

    document.add_heading("Next Recommended Experiment", level=1)
    experiment = payload.get("next_recommended_experiment")
    if experiment:
        document.add_paragraph(experiment["objective"])
        variables = experiment["variables"]
        document.add_paragraph(
            "Variables: "
            f"T={variables['temperature']:.2f} K, residence_time={variables['residence_time']:.2f}, "
            f"light={variables['light_intensity']:.2f}, catalyst={variables['catalyst_loading']:.3f}"
        )
    else:
        document.add_paragraph("No next experiment has been proposed.")

    document.add_heading("Uncertainty Analysis", level=1)
    for ranking in payload["uncertainty_analysis"][:5]:
        document.add_paragraph(
            f"{ranking['hypothesis_id']}: uncertainty={ranking['uncertainty']:.3f}; "
            f"caveats={'; '.join(ranking.get('caveats', [])[:2])}",
            style="List Bullet",
        )

    document.add_heading("DFT Jobs Submitted", level=1)
    if payload["dft_jobs_submitted"]:
        for job in payload["dft_jobs_submitted"]:
            document.add_paragraph(
                f"{job['id']}: {job['calculation_type']} for {job['linked_hypothesis_id']}",
                style="List Bullet",
            )
    else:
        document.add_paragraph("No DFT jobs submitted.")

    document.add_heading("Robot Experiments Submitted", level=1)
    if payload["robot_experiments_submitted"]:
        for experiment in payload["robot_experiments_submitted"]:
            document.add_paragraph(experiment["id"], style="List Bullet")
    else:
        document.add_paragraph("No robot experiments submitted.")

    document.add_heading("Critic Notes", level=1)
    if payload["critic_notes"]:
        for note in payload["critic_notes"][:12]:
            document.add_paragraph(note, style="List Bullet")
    else:
        document.add_paragraph("No critic notes recorded.")

    if payload["errors"]:
        document.add_heading("Errors", level=1)
        for error in payload["errors"]:
            document.add_paragraph(error, style="List Bullet")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def _write_minimal_docx(payload: dict, output_path: Path, *, reason: str) -> None:
    """Write a simple valid DOCX when python-docx is unavailable."""

    paragraphs = [
        payload["title"],
        f"Objective: {payload['objective']}",
        f"Mode: {payload['mode']}",
        f"Round: {payload['round_index']}",
        "Current Best Mechanism Ranking",
    ]
    for ranking in payload.get("best_mechanism_ranking", [])[:5]:
        paragraphs.append(
            f"{ranking['hypothesis_id']} | {ranking['mechanism_class']} | "
            f"score={ranking['score']:.3f} | uncertainty={ranking['uncertainty']:.3f}"
        )
    experiment = payload.get("next_recommended_experiment")
    if experiment:
        paragraphs.append(f"Next Recommended Experiment: {experiment['objective']}")
    paragraphs.append("Critic Notes")
    paragraphs.extend(payload.get("critic_notes", [])[:12])
    if payload.get("errors"):
        paragraphs.append("Errors")
        paragraphs.extend(payload["errors"])
    paragraphs.append(f"Fallback DOCX writer used because python-docx was unavailable: {reason}")

    body = "".join(f"<w:p><w:r><w:t>{escape(str(paragraph))}</w:t></w:r></w:p>" for paragraph in paragraphs)
    document_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f"<w:body>{body}<w:sectPr /></w:body></w:document>"
    )
    content_types = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/word/document.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
        "</Types>"
    )
    rels = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
        'Target="word/document.xml"/>'
        "</Relationships>"
    )
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(output_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", content_types)
        archive.writestr("_rels/.rels", rels)
        archive.writestr("word/document.xml", document_xml)
